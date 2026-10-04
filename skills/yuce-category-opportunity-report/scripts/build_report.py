#!/usr/bin/env python3
"""Prepare a local Excel cache and render a Yuce category research report."""

import argparse
import hashlib
import html
import json
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from market_routes import build_routes, LABELS

try:
    from openpyxl import Workbook, load_workbook
except ImportError as exc:
    raise SystemExit("缺少 openpyxl；请在稳定 Python 环境安装 openpyxl>=3.1 后重试") from exc

CONTRACT = "yuce-category-opportunity@2.0"
CACHE_SHEET = "类目月度缓存"
CACHE_HEADERS = ["月份", "层级", "一级类目", "二级类目", "三级类目", "成交额（原始值）",
                 "成交额环比（原始值）", "成交量（原始值）", "成交量环比（原始值）",
                 "是否有三级类目", "来源文件", "来源SHA256"]
SOURCE_SHEETS = {"一级类目月数据": 1, "二级类目月数据": 2, "三级类目月数据": 3}


def months_between(start, end):
    try:
        year, month = (int(part) for part in start.split("-"))
        end_year, end_month = (int(part) for part in end.split("-"))
        date(year, month, 1)
        date(end_year, end_month, 1)
    except (ValueError, TypeError) as exc:
        raise ValueError("月份格式必须为 YYYY-MM") from exc
    months = []
    while (year, month) <= (end_year, end_month) and len(months) <= 120:
        months.append(f"{year:04d}-{month:02d}")
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    if not months or months[-1] != end or len(months) != 24:
        raise ValueError("报告需要连续 24 个月的数据窗口")
    return months


def read_json(path):
    raw = Path(path).read_bytes()
    return json.loads(raw), hashlib.sha256(raw).hexdigest()


def decimal_amount(value):
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, TypeError) as exc:
        raise ValueError(f"成交额无效：{value}") from exc
    if not amount.is_finite() or amount < 0:
        raise ValueError(f"成交额无效：{value}")
    return amount


def normalize_month(value):
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m")
    if isinstance(value, str) and len(value) == 7 and value[4] == "-":
        try:
            date(int(value[:4]), int(value[5:]), 1)
            return value
        except ValueError:
            pass
    raise ValueError(f"月份无效：{value}")


def normalize_yc(data, source):
    kind = data.get("kind")
    level = {"first-category-monthly": 1, "second-category-monthly": 2,
             "third-category-monthly": 3}.get(kind)
    if not level or not isinstance(data.get("rows"), list):
        raise ValueError(f"不是 yccli 类目月数据 JSON：{source}")
    rows = []
    for item in data["rows"]:
        names = [item.get("firstCategory")]
        if level >= 2:
            names.append(item.get("secondCategory"))
        if level >= 3:
            names.append(item.get("thirdCategory"))
        if any(not isinstance(name, str) or not name.strip() for name in names):
            raise ValueError(f"类目路径缺失：{source}")
        rows.append({"month": normalize_month(item.get("month")), "level": level, "path": names,
                     "amount": decimal_amount(item.get("transactionAmount")),
                     "amountMom": item.get("transactionAmountMom"),
                     "volume": item.get("transactionVolume") if level == 1 else None,
                     "volumeMom": item.get("transactionVolumeMom") if level == 1 else None,
                     "hasThirdCategory": item.get("hasThirdCategory") if level == 2 else None,
                     "source": source})
    return rows


def normalize_db(data, source):
    if data.get("kind") != "yuce-category-monthly-snapshot":
        raise ValueError(f"不是 tbcli 预策快照：{source}")
    return [{"month": normalize_month(item["month"]), "level": int(item["level"]), "path": item["path"],
             "amount": decimal_amount(item["transactionAmount"]),
             "amountMom": item.get("transactionAmountMom"),
             "volume": item.get("transactionVolume"),
             "volumeMom": item.get("transactionVolumeMom"),
             "hasThirdCategory": item.get("hasThirdCategory"), "source": source}
            for item in data.get("rows", [])]


def _value(record, headers, name):
    return record[headers.index(name)] if name in headers else None


def normalize_excel(source, category, cache=False):
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows = []
    matched = 0
    try:
        if cache and workbook.sheetnames != [CACHE_SHEET]:
            raise ValueError(f"本地缓存必须只有 {CACHE_SHEET} 数据表：{source}")
        for sheet in workbook:
            level = 0 if cache and sheet.title == CACHE_SHEET else SOURCE_SHEETS.get(sheet.title)
            if level is None:
                continue
            matched += 1
            values = sheet.values
            try:
                headers = [str(value).strip() if value is not None else "" for value in next(values)]
            except StopIteration as exc:
                raise ValueError(f"Excel 数据表为空：{source} / {sheet.title}") from exc
            required = CACHE_HEADERS if level == 0 else (["月份", "一级类目", "成交额（原始值）"] if level == 1 else
                       ["月份", "一级类目", "二级类目", "总成交额（原始值）"] + (["三级类目"] if level == 3 else []))
            missing = [name for name in required if name not in headers]
            if missing:
                raise ValueError(f"Excel 缺少字段 {missing}：{source} / {sheet.title}")
            for line, record in enumerate(values, start=2):
                if not any(value is not None and value != "" for value in record):
                    continue
                actual_level = int(_value(record, headers, "层级")) if level == 0 else level
                if actual_level not in (1, 2, 3):
                    raise ValueError(f"Excel 层级无效：{source} / {sheet.title} 第{line}行")
                names = [_value(record, headers, "一级类目")]
                if actual_level >= 2:
                    names.append(_value(record, headers, "二级类目"))
                if actual_level >= 3:
                    names.append(_value(record, headers, "三级类目"))
                if any(not isinstance(name, str) or not name.strip() for name in names) or names[0] != category:
                    raise ValueError(f"Excel 类目路径不匹配：{source} / {sheet.title} 第{line}行")
                amount_name = "成交额（原始值）" if actual_level == 1 or level == 0 else "总成交额（原始值）"
                mom_name = "成交额环比（原始值）"
                rows.append({"month": normalize_month(_value(record, headers, "月份")), "level": actual_level, "path": names,
                             "amount": decimal_amount(_value(record, headers, amount_name)),
                             "amountMom": _value(record, headers, mom_name),
                             "volume": _value(record, headers, "成交量（原始值）") if actual_level == 1 else None,
                             "volumeMom": _value(record, headers, "成交量环比（原始值）") if actual_level == 1 else None,
                             "hasThirdCategory": _value(record, headers, "是否有三级类目") if actual_level == 2 else None,
                             "source": _value(record, headers, "来源文件") if level == 0 else source,
                             "sourceSha256": _value(record, headers, "来源SHA256") if level == 0 else None})
                if level == 0 and (not rows[-1]["source"] or not isinstance(rows[-1]["sourceSha256"], str)
                                   or len(rows[-1]["sourceSha256"]) != 64):
                    raise ValueError(f"缓存来源身份缺失：{source} / {sheet.title} 第{line}行")
    finally:
        workbook.close()
    if matched == 0 or not rows:
        raise ValueError(f"Excel 没有可识别的类目月数据表：{source}")
    return rows


def merge_rows(category, months, db_path=None, yc_paths=(), excel_paths=(), cache_path=None):
    if cache_path and (db_path or yc_paths or excel_paths):
        raise ValueError("本地 Excel 缓存不能与原始来源混用")
    paths = ([cache_path] if cache_path else []) + ([db_path] if db_path else []) + list(yc_paths) + list(excel_paths)
    if not paths:
        raise ValueError("至少需要一个本地 Excel 缓存、原始 Excel、tbcli 快照或 yccli JSON")
    merged = {}
    provenance = []
    for source in paths:
        digest = hashlib.sha256(Path(source).read_bytes()).hexdigest()
        if source == cache_path or source in excel_paths:
            rows = normalize_excel(source, category, cache=(source == cache_path))
            kind = "local-xlsx-cache" if source == cache_path else "source-xlsx"
        else:
            data, _ = read_json(source)
            rows = normalize_db(data, source) if source == db_path else normalize_yc(data, source)
            if data.get("category") != category:
                raise ValueError(f"来源的一级类目不匹配：{source}")
            kind = data.get("kind")
        provenance.append({"file": Path(source).name, "sha256": digest, "kind": kind})
        for row in rows:
            if row["month"] not in months:
                continue
            if len(row["path"]) != row["level"] or row["path"][0] != category:
                raise ValueError(f"来源的类目层级不匹配：{source}")
            row["sourceSha256"] = row.get("sourceSha256") or digest
            key = (row["level"], tuple(row["path"]), row["month"])
            prior = merged.get(key)
            if prior:
                if abs(prior["amount"] - row["amount"]) > Decimal("0.01"):
                    raise ValueError(f"同一类目月份的来源数值冲突：{row['path']} {row['month']}")
                continue
            merged[key] = row
            if source == cache_path and row.get("source") and row.get("sourceSha256"):
                item = {"file": Path(str(row["source"])).name, "sha256": row["sourceSha256"], "kind": "cache-origin"}
                if item not in provenance:
                    provenance.append(item)
    if not merged:
        raise ValueError("选定窗口没有可分析的类目月数据")
    return merged, provenance


def write_cache(output, merged):
    def optional_number(value):
        if value is None or value == "":
            return None
        try:
            parsed = Decimal(str(value))
        except (InvalidOperation, TypeError) as exc:
            raise ValueError(f"可选数值字段无效：{value}") from exc
        if not parsed.is_finite():
            raise ValueError(f"可选数值字段无效：{value}")
        return float(parsed)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = CACHE_SHEET
    sheet.append(CACHE_HEADERS)
    for (_, path, month), row in sorted(merged.items(), key=lambda item: (item[0][2], item[0][0], item[0][1])):
        sheet.append([month, row["level"], path[0], path[1] if len(path) > 1 else None,
                      path[2] if len(path) > 2 else None, float(row["amount"]), optional_number(row.get("amountMom")),
                      optional_number(row.get("volume")), optional_number(row.get("volumeMom")), row.get("hasThirdCategory"),
                      Path(str(row["source"])).name, row.get("sourceSha256")])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 24
    sheet.column_dimensions["E"].width = 48
    sheet.column_dimensions["K"].width = 38
    for cell in sheet["F"][1:]:
        cell.number_format = "#,##0.00"
    with output.open("xb") as target:
        workbook.save(target)
    workbook.close()


def compute_stats(category, months, merged):
    grouped = defaultdict(dict)
    for (level, path, month), row in merged.items():
        grouped[(level, path)][month] = row
    root_key = (1, (category,))
    if root_key not in grouped or any(month not in grouped[root_key] for month in months):
        raise ValueError("一级类目缺少完整的连续 24 个月，不能生成同比报告")
    stats = []
    for (level, path), series in grouped.items():
        prior = months[:12]
        recent = months[12:]
        complete = all(month in series for month in months)
        before = sum((series[month]["amount"] for month in prior if month in series), Decimal(0))
        after = sum((series[month]["amount"] for month in recent if month in series), Decimal(0))
        yoy = float(after / before - 1) if complete and before > 0 else None
        recent6_before = sum((series[month]["amount"] for month in prior[6:] if month in series), Decimal(0))
        recent6_after = sum((series[month]["amount"] for month in recent[6:] if month in series), Decimal(0))
        six_yoy = float(recent6_after / recent6_before - 1) if complete and recent6_before > 0 else None
        positive_months = sum(series[recent[index]]["amount"] > series[prior[index]]["amount"]
                              for index in range(12)) if complete else None
        stats.append({"level": level, "path": list(path), "monthCount": len(series),
                      "missingMonths": [month for month in months if month not in series],
                      "monthlyAmounts": [float(series[month]["amount"]) if month in series else None for month in months],
                      "priorAmount": float(before), "recentAmount": float(after),
                      "increment": float(after - before) if complete else None,
                      "yoy": yoy, "sixMonthYoy": six_yoy,
                      "positiveYoyMonths": positive_months, "complete": complete})
    stats.sort(key=lambda item: (item["level"], -item["recentAmount"], item["path"]))
    root = next(item for item in stats if item["level"] == 1)
    for item in stats:
        item["rootShare"] = item["recentAmount"] / root["recentAmount"] if root["recentAmount"] else None
    return stats


def reconcile(months, merged):
    categories = defaultdict(dict)
    for (level, path, month), row in merged.items():
        categories[(level, path)][month] = row["amount"]
    issues = []
    for (level, path), parent_series in categories.items():
        if level == 3:
            continue
        children = [(child_path, series) for (child_level, child_path), series in categories.items()
                    if child_level == level + 1 and child_path[:-1] == path]
        if not children:
            continue
        for month in months:
            if month not in parent_series or any(month not in series for _, series in children):
                continue
            difference = parent_series[month] - sum((series[month] for _, series in children), Decimal(0))
            tolerance = max(Decimal("1.00"), Decimal("0.01") * len(children))
            if abs(difference) > tolerance:
                issues.append({"path": list(path), "month": month, "difference": float(difference)})
    return issues


def quantitative_watch(stats, candidates, benchmark, screen):
    excluded = {tuple(item["path"]) for item in candidates}
    if benchmark:
        excluded.add(tuple(benchmark["path"]))
    possible = [item for item in stats if item["level"] == 3 and item["complete"]
                and item["recentAmount"] >= screen["minRecentAmount"]
                and tuple(item["path"]) not in excluded
                and item["yoy"] is not None
                and (item["yoy"] <= 0 or item["sixMonthYoy"] is not None
                     and item["sixMonthYoy"] < item["yoy"] - 0.15)]
    possible.sort(key=lambda item: (-item["recentAmount"], item["path"]))
    for item in possible[:5]:
        item["watchReason"] = (f'最近12个月成交额 {money(item["recentAmount"])}，全年同比 {percent(item["yoy"])}，'
                               f'近6个月同比 {percent(item["sixMonthYoy"])}；'
                               + ("全年成交额下降，先观察是否延续。" if item["yoy"] <= 0 else
                                  "近期同比比全年明显偏低，先核对增长是否放缓。"))
    return possible[:5]


def money(value):
    if value is None:
        return "—"
    if abs(value) >= 1e8:
        return f"{value / 1e8:.2f}亿"
    if abs(value) >= 1e4:
        return f"{value / 1e4:.2f}万"
    return f"{value:,.0f}"


def percent(value):
    return "—" if value is None else f"{value:+.1%}"


def index_points(value):
    return "—" if value is None else f"{value:.2f} 点"


def esc(value):
    return html.escape(str(value), quote=True)


def card(title, body):
    return f'<article class="card"><div class="card-head"><h2>{esc(title)}</h2></div><div class="card-body">{body}</div></article>'


def chart_attribute(spec):
    for group in spec["groups"]:
        for series in group["series"]:
            if len(series["values"]) != len(spec["labels"]):
                raise ValueError("图表与表格行数不一致")
    return ' data-chart-spec="' + esc(json.dumps(spec, ensure_ascii=False, allow_nan=False)) + '"'


def stats_chart(items, keys):
    definitions = {
        "priorAmount": ("成交规模", "元", "前12个月"), "recentAmount": ("成交规模", "元", "最近12个月"),
        "yoy": ("同比增速", "%", "全年同比"), "sixMonthYoy": ("同比增速", "%", "近6个月同比"),
        "increment": ("成交增量", "元", "年度增量"), "positiveYoyMonths": ("增长月份", "个月", "增长月份"),
        "growthImpactIndex": ("持续增量贡献", "点", "持续增量贡献"),
        "sustainedGrowthIndex": ("持续增长强度", "点", "持续增长强度"),
        "capacityPercentile": ("容量百分位", "百分位", "容量百分位"),
        "routeRank": ("路线内排名（越小越靠前）", "名", "路线内排名"),
        "growthAccelerationPp": ("半年增速差", "个百分点", "后半年减前半年同比"),
        "top3PositiveIncrementShare": ("正增量前三月占比", "%", "前三月正增量占比"),
        "rootShare": ("近期占比", "%", "一级类目GMV占比"),
    }
    groups = {}
    for key in keys:
        label, unit, name = definitions[key]
        group = groups.setdefault(label, {"label": label, "unit": unit, "series": []})
        group["series"].append({"name": name, "values": [x.get(key) if x["complete"] else None for x in items]})
    return {"kind": "bar", "labels": [x["path"][-1] for x in items], "groups": list(groups.values())}


def table(title, headers, data_rows, chart):
    head = "".join(f"<th>{esc(header)}</th>" for header in headers)
    body = "".join("<tr class=\"searchable\">" + "".join(f"<td>{esc(value)}</td>" for value in row) + "</tr>" for row in data_rows)
    if len(chart["labels"]) != len(data_rows):
        raise ValueError("图表数据必须覆盖表格所有行")
    return f'<article class="card"><div class="card-head"><h2>{esc(title)}</h2><span class="muted">{len(data_rows)} 行</span></div><div class="table-wrap"><table class="data-table"{chart_attribute(chart)}><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div></article>'


def view(view_id, eyebrow, title, subtitle, body, badge="", show_heading=True):
    if not show_heading:
        return f'<section class="view" data-view="{esc(view_id)}" id="view-{esc(view_id)}" tabindex="-1" hidden>{body}</section>'
    heading = f'<div class="page-heading"><div><div class="eyebrow">{esc(eyebrow)}</div><h1>{esc(title)}</h1><p>{esc(subtitle)}</p></div>'
    if badge:
        heading += f'<span class="tag">{esc(badge)}</span>'
    heading += "</div>"
    return f'<section class="view" data-view="{esc(view_id)}" id="view-{esc(view_id)}" tabindex="-1" hidden>{heading}{body}</section>'


def comparison_rows(item, months):
    values = item["monthlyAmounts"]
    rows = []
    for index in range(12):
        before, after = values[index], values[index + 12]
        change = after / before - 1 if before is not None and before > 0 and after is not None else None
        rows.append((months[index][5:], money(before), money(after), percent(change)))
    return rows


def numbered_insights(title, items):
    return (f'<article class="card insight-card"><div class="card-head"><h2>{esc(title)}</h2></div>'
            '<ol class="insight-list">' + ''.join(f'<li><span class="insight-number">{index}</span>'
            f'<div class="insight-content">{content}</div></li>' for index, content in enumerate(items, 1))
            + '</ol></article>')


def growth_tempo(item):
    annual, recent = item["yoy"], item["sixMonthYoy"]
    if annual is None or recent is None:
        return "最近6个月无法形成完整同月比较。"
    difference = recent - annual
    if difference > 0.03:
        return f'近6个月同比比全年高 <strong class="positive-text">{difference * 100:.1f} 个百分点</strong>，近期增速高于全年。'
    if difference < -0.03:
        return f'近6个月同比比全年低 <strong class="warning-text">{-difference * 100:.1f} 个百分点</strong>，近期增速低于全年。'
    return f'近6个月同比与全年相差 <strong>{abs(difference) * 100:.1f} 个百分点</strong>，两者接近。'


def candidate_detail(item, months, benchmark=False):
    name = item["path"][-1]
    label = "规模基准" if benchmark else "优先深研"
    verify = item.get("nextVerification") or "核查商品和店铺排行、价格带、评价、供应链与完整经营成本。"
    rank = item.get("routeRank")
    rank_text = f'{item.get("capacityLabel", "")} · 路线内第 {rank} 位' if rank else "规模参照；未入本轮研究池"
    metric = "growthImpactIndex" if item.get("capacityBand") == "large" else "sustainedGrowthIndex"
    metric_label = "持续增量贡献" if metric == "growthImpactIndex" else "持续增长强度"
    observations = [
        (f'<strong>体量与贡献：</strong>最近12个月成交额 <strong class="info-text">{esc(money(item["recentAmount"]))}</strong>，'
         f'占一级类目 {percent(item["rootShare"])}；前12个月为 {esc(money(item["priorAmount"]))}，'
         f'实际新增 {esc(money(item["increment"]))}。'),
        (f'<strong>增长与排序：</strong>全年同比 <strong class="positive-text">{esc(percent(item["yoy"]))}</strong>，'
         f'{metric_label} <strong class="info-text">{esc(index_points(item.get(metric)))}</strong>；'
         f'{esc(rank_text)}。指标仅决定各自路线内的研究顺序，不衡量竞争强弱。'),
        (f'<strong>月度持续性：</strong>12组同月中 <strong>{item["positiveYoyMonths"]}/12 组增长</strong>；'
         f'近6个月同比 {esc(percent(item["sixMonthYoy"]))}。{growth_tempo(item)}'
         f' 最近一年前6个月同比 {esc(percent(item.get("firstHalfYoy")))}；'
         f'增量最大的3个月占全年正向月增量 {esc(percent(item.get("top3PositiveIncrementShare")))}。'
         '集中度高可能与季节性有关，须核对逐月证据。'),
    ]
    body = (f'<div class="card-body"><span class="tag">{label}</span><p class="prose">{esc(" / ".join(item["path"][1:]))}</p>'
            f'<div class="detail-metrics"><div><span>最近12个月</span><b>{esc(money(item["recentAmount"]))}</b></div>'
            f'<div><span>全年同比</span><b>{esc(percent(item["yoy"]))}</b><small>{item["positiveYoyMonths"]}/12 个月增长</small></div>'
            f'<div><span>成交增量</span><b>{esc(money(item["increment"]))}</b></div>'
            f'<div><span>{metric_label}</span><b>{esc(index_points(item.get(metric)))}</b></div></div></div>')
    return (f'<article class="card candidate-detail searchable"><div class="card-head"><h2>{esc(name)}</h2></div>{body}'
            + numbered_insights("数据解读", observations)
            + table("12组同月成交额对比", ["对应月份", "前一年", "最近一年", "同比"], comparison_rows(item, months),
                    {"kind": "line", "labels": [m[5:] + "月" for m in months[:12]], "groups": [
                        {"label": "对应月份成交额", "unit": "元", "series": [
                            {"name": f"前一年 {months[0]}—{months[11]}", "values": item["monthlyAmounts"][:12]},
                            {"name": f"最近一年 {months[12]}—{months[23]}", "values": item["monthlyAmounts"][12:]}]},
                        {"label": "同月同比", "unit": "%", "series": [{"name": "同月同比", "values": [
                            after / before - 1 if before is not None and before > 0 and after is not None else None
                            for before, after in zip(item["monthlyAmounts"][:12], item["monthlyAmounts"][12:])]}]}]})
            + f'<div class="card-body verification"><h3>下一步需要补的证据</h3><p class="prose">{esc(verify)}</p>'
              '<p class="prose muted">类目月度数据不能证明具体使用场景、产品差异或白牌利润。</p></div>'
            + '</article>')


def route_cards(model, stats, candidates, compact=False):
    lookup = {tuple(x["path"]): x for x in stats if x["level"] == 3}
    ids = {tuple(x["path"]): f"candidate-{i}" for i, x in enumerate(candidates, 1)}
    cuts = model["boundaryAmounts"]
    ranges = {"large": f'≥ {money(cuts[2])}', "medium": f'{money(cuts[1])} ≤ GMV < {money(cuts[2])}',
              "small": f'{money(cuts[0])} ≤ GMV < {money(cuts[1])}'}
    if cuts[0] is None:
        ranges = {key: "无完整正成交总体，分界不可计算" for key in ranges}
    sections = []
    for route in model["routes"]:
        rows = []
        paths = route["candidatePaths"][:3] if compact else route["candidatePaths"]
        for path in paths:
            x = lookup[tuple(path)]
            tempo = x.get("growthAccelerationPp")
            reading = ("前后半年增速接近" if tempo is not None and abs(tempo) <= 3 else
                       "后半年增速较高" if tempo is not None and tempo > 3 else
                       "后半年增速较低" if tempo is not None else "半年基期不足")
            rows.append(f'<tr class="searchable"><td>{x["routeRank"]}</td><td><button class="read-row-link" data-view-target="{ids[tuple(path)]}">{esc(path[-1].split(" > ")[-1])}</button></td>'
                        f'<td>{money(x["recentAmount"])}</td><td>{percent(x["yoy"])}</td><td>{percent(x["sixMonthYoy"])}</td>'
                        f'<td>{x["positiveYoyMonths"]}/12</td><td>{index_points(x[route["metric"]])}</td><td>{reading}</td></tr>')
        body = (f'<p class="prose">本档 <strong>{route["bandCount"]}</strong> 个 → 增长入池 <strong>{route["poolCount"]}</strong> 个 → 首批 <strong>{len(route["candidatePaths"])}</strong> 个。'
                f'实际年GMV边界：<strong>{esc(ranges[route["id"]])}</strong>。</p>')
        body += ('<p class="prose">按持续增量贡献排序，关注绝对新增需求。</p>' if route["id"] == "large" else
                 '<p class="prose">按持续增长强度排序；容量达标后不再为规模加分。</p>')
        if rows:
            spec = stats_chart([lookup[tuple(path)] for path in paths], ["recentAmount", "yoy", "sixMonthYoy", "positiveYoyMonths", route["metric"], "routeRank"])
            body += (f'<div class="table-wrap"><table class="data-table"{chart_attribute(spec)}><thead><tr><th>路线内顺序</th><th>三级类目</th><th>近12个月</th><th>全年同比</th><th>近6个月同比</th><th>增长月份</th>'
                     f'<th>{route["metricLabel"]}·点</th><th>数据解读</th></tr></thead><tbody>' + ''.join(rows) + '</tbody></table></div>')
        else:
            body += f'<p class="notice">{esc(route["emptyReason"])}</p>'
        if compact:
            body += '<p class="prose muted">首页每条路线最多预览前三名。</p><button class="read-row-link" data-view-target="candidates">查看各路线完整首批名单与逐类目证据 →</button>'
        sections.append(card(route["label"], body))
    return ''.join(sections)


def route_method(model):
    qs, cuts, dist = model["policy"]["capacityQuantiles"], model["boundaryAmounts"], model["distribution"]
    p = model["policy"]
    text = (f'<p class="prose">分组总体：本一级类目中24个月完整且近一年成交额为正的三级类目，共 <strong>{dist["populationCount"]}</strong> 个；'
            f'中位数 <strong>{money(dist["median"])}</strong>；按数量取规模前10%（向上取整）的类目，占分组总体成交额 <strong>{percent(dist["top10PercentGmvShare"])}</strong>。</p>'
            f'<p class="prose">本次分界：' + '；'.join(f'P{q*100:g} = {money(c)}' for q, c in zip(qs, cuts)) + '。分位数采用线性插值，金额仅是计算结果，分档使用未四舍五入值。</p>'
            f'<p class="prose">低于P{qs[0]*100:g}归入相对小容量观察；其上依次为中小、中等、大容量。同额类目保持同档，边界值归较高档。分位点是可调整研究参数，不是行业通用真理。</p>'
            f'<p class="prose">选择理由：{esc(p["why"])}。全量表“容量百分位”采用同额中秩，描述相对位置，与线性插值边界口径不同。</p>'
            f'<p class="prose">增长条件：全年同比≥{percent(p["minYoy"])}，近6个月同比≥{percent(p["minSixMonthYoy"])}，增长月份≥{p["minPositiveMonths"]}/12；每条路线最多{p["shortlistLimit"]}个。混合“其他”类目先拆解再研究。</p>'
            '<p class="prose">大市场：持续增量贡献＝100 × max(年度增量,0) ÷ 一级类目年GMV × 增长月份/12。中等与中小：持续增长强度＝100 × min(全年同比,近6个月同比) × 增长月份/12。两种指标不跨路线比较；均不预测利润。全年与半年窗口重叠，并非两份独立证据。</p>'
            '<p class="prose">加速信号：最近一年后6个月同比减前6个月同比；相差3个百分点以内记为接近。正向月增量集中度：最大3个月正增量 ÷ 全年各月正增量之和；用于检查季节性和波动，不直接淘汰。</p>')
    for s in model["sensitivity"]:
        text += f'<p class="prose">敏感性试算：分位点改为 {esc(str([round(q, 3) for q in s["quantiles"]]))}，有 {s["changedBandCount"]} 个类目改变容量档位。</p>'
    text += '<p class="prose"><em>相对中小不代表绝对容量足够。仍需结合目标销售额评估所需市场份额；品牌集中度和白牌成交情况尚未验证。</em></p>'
    text += ''.join(f'<p class="notice">{esc(w)}</p>' for w in model["warnings"])
    return card("容量如何随类目变化", text)


def multi_route_overview(category, months, stats, model, candidates):
    root = next(x for x in stats if x["level"] == 1)
    branches = [x for x in stats if x["level"] == 2 and x["complete"]]
    lead = max(branches, key=lambda x: x["increment"], default=None)
    lead_text = (f'{esc(lead["path"][-1])}的年度增量为 {money(lead["increment"])}。' if lead else '二级类目缺少完整可比数据。')
    if lead and root["increment"] > 0:
        lead_text += f'占一级类目总增量 {lead["increment"]/root["increment"]:.1%}。'
    hero = (f'<div class="read-hero"><div class="read-overline">{esc(category)} · {esc(months[0])}—{esc(months[-1])}</div>'
            '<h1>先选适合自己的市场，再验证白牌机会</h1><p>三条路线同时保留，分别排序；中小市场不因规模较小就被判定竞争更弱。</p></div>')
    hero += numbered_insights("这份数据最重要的三句话", [
        f'<strong>先看市场：</strong>近12个月成交额 <strong class="info-text">{money(root["recentAmount"])}</strong>，全年同比 <strong>{percent(root["yoy"])}</strong>，近6个月同比 <strong>{percent(root["sixMonthYoy"])}</strong>。',
        f'<strong>再看结构：</strong>{lead_text}各容量路线都在同一个一级类目内比较，不把二级分支各自的小市场误称为全局中等市场。',
        '<strong>最后选路线：</strong>大市场看新增成交量级，中等和中小市场看持续增长；下一步先查品牌集中度、非头部商家的成交表现，再投入深度商品研究。',
    ])
    hero += route_cards(model, stats, candidates, compact=True)
    hero += '<details class="card"><summary class="card-body">展开容量分组、门槛与指标依据</summary>' + route_method(model) + '</details>'
    hero += card("先研究哪里，不等于现在卖什么", '<ol class="list"><li>轻量核查商品与店铺榜单：品牌、商家集中度、非头部品牌成交；注明样本覆盖，不能直接外推全市场。</li><li>验证搜索需求、评价、商品形态与差异化空间。</li><li>结合目标销售额、供应链及完整成本决定是否试做。</li></ol>')
    return hero


def render_html(category, months, stats, pool, candidates, benchmark, watch, screen, issues, provenance, template, decision_mode, route_model):
    root = next(item for item in stats if item["level"] == 1)
    incomplete = [item for item in stats if not item["complete"]]
    overview = multi_route_overview(category, months, stats, route_model, candidates)
    structure_rows = [(item["path"][-1], money(item["priorAmount"]), money(item["recentAmount"]),
                       percent(item["yoy"]), money(item["increment"]),
                       f'{item["increment"] / root["increment"]:.1%}' if root["increment"] > 0 and item["complete"] else "—",
                       percent(item["rootShare"]),
                       f'{(item["rootShare"] - item["priorAmount"] / root["priorAmount"]) * 100:+.1f}个百分点' if root["priorAmount"] and root["recentAmount"] and item["complete"] else "—",
                       f'{item["positiveYoyMonths"]}/12')
                      for item in stats if item["level"] == 2]
    top_branch = max((item for item in stats if item["level"] == 2 and item["complete"]), key=lambda item: item["increment"], default=None)
    branches = [item for item in stats if item["level"] == 2 and item["complete"]]
    branch_count = sum(item["yoy"] is not None and item["yoy"] > 0 for item in branches)
    biggest_branch = max(branches, key=lambda item: item["recentAmount"], default=None)
    gain_branch = max(branches, key=lambda item: item["rootShare"] - item["priorAmount"] / root["priorAmount"], default=None) if root["priorAmount"] and root["recentAmount"] else None
    loss_branch = min(branches, key=lambda item: item["rootShare"] - item["priorAmount"] / root["priorAmount"], default=None) if root["priorAmount"] and root["recentAmount"] else None
    share_reading = (f'<strong>份额再分配：</strong>当前规模最大的二级类目是 <strong class="info-text">{esc(biggest_branch["path"][-1])}</strong>，'
                     f'占一级类目 {biggest_branch["rootShare"]:.1%}。'
                     f'{esc(gain_branch["path"][-1])}的份额变化最大（{(gain_branch["rootShare"] - gain_branch["priorAmount"] / root["priorAmount"]) * 100:+.1f} 个百分点）；'
                     f'{esc(loss_branch["path"][-1])}的份额变化最小（{(loss_branch["rootShare"] - loss_branch["priorAmount"] / root["priorAmount"]) * 100:+.1f} 个百分点）。'
                     if biggest_branch and gain_branch and loss_branch else
                     '<strong>份额再分配：</strong>数据不足以计算二级类目的可比份额变化。')
    structure_insights = [
        f'<strong>一级总体：</strong>最近12个月成交额 <strong class="info-text">{esc(money(root["recentAmount"]))}</strong>，相对前12个月增长 <strong class="positive-text">{esc(percent(root["yoy"]))}</strong>，实际新增 {esc(money(root["increment"]))}。',
        (f'<strong>增量来源：</strong><strong class="info-text">{esc(top_branch["path"][-1])}</strong>新增 {esc(money(top_branch["increment"]))}，'
         f'占一级类目总增量 <strong>{top_branch["increment"] / root["increment"]:.1%}</strong>；'
         f'其最近12个月一级占比为 {top_branch["rootShare"]:.1%}。' if top_branch and root["increment"] > 0 else
         '<strong>增量来源：</strong>一级类目无正向总增量或缺少完整二级类目，不计算正向增量贡献率。'),
        f'<strong>增长广度：</strong>{len(branches)} 个数据完整的二级类目中，<strong>{branch_count} 个</strong>全年同比为正。下表展示规模、增量与同月增长数；逐月表用于查看变化是否集中在少数月份。',
        share_reading,
        f'<strong>近期速度：</strong>一级类目近6个月同月同比 <strong>{esc(percent(root["sixMonthYoy"]))}</strong>，全年同比 {esc(percent(root["yoy"]))}。{growth_tempo(root)}',
    ]
    structure = numbered_insights("市场结构的数据解读", structure_insights)
    structure_items = [x for x in stats if x["level"] == 2]
    spec = stats_chart(structure_items, ["priorAmount", "recentAmount", "yoy", "increment", "rootShare", "positiveYoyMonths"])
    spec["groups"] += [
        {"label": "增量贡献", "unit": "%", "series": [{"name": "增量贡献", "values": [x["increment"] / root["increment"] if x["complete"] and root["increment"] > 0 else None for x in structure_items]}]},
        {"label": "占比变化", "unit": "个百分点", "series": [{"name": "占比变化", "values": [(x["rootShare"] - x["priorAmount"] / root["priorAmount"]) * 100 if x["complete"] and root["priorAmount"] and root["recentAmount"] else None for x in structure_items]}]}]
    structure += table(f"{len(branches)} 个二级类目的规模与增量", ["二级类目", "前12个月", "最近12个月", "同比", "成交增量", "增量贡献", "近期占比", "占比变化", "增长月份"], structure_rows, spec)
    second = {item["path"][-1]: item for item in stats if item["level"] == 2}
    branch_names = list(second)
    monthly_rows = [(month, money(root["monthlyAmounts"][index])) + tuple(money(second[name]["monthlyAmounts"][index]) for name in branch_names)
                    for index, month in enumerate(months)]
    structure += table("一级及二级类目逐月成交额", ["月份", category] + branch_names, monthly_rows,
                       {"kind": "line", "labels": months, "groups": [{"label": "月度成交额", "unit": "元", "series": [
                           {"name": x["path"][-1], "values": x["monthlyAmounts"]} for x in [root] + list(second.values())]}]})
    structure += card("层级对账", f'<p class="prose">可计算的父子月份中，超过1元容差的差异：<strong>{len(issues)}</strong> 处。类目金额按层级分别观察，不把一级、二级和三级相加。</p>')
    candidate = route_cards(route_model, stats, candidates)
    watch_rows = [(" / ".join(item["path"][1:]), money(item["recentAmount"]), percent(item["yoy"]),
                   percent(item["sixMonthYoy"]), item["watchReason"]) for item in watch]
    watch_page = table("需要谨慎判断的类目", ["三级类目", "近期成交额", "全年同比", "近6个月同比", "处理理由"], watch_rows,
                       stats_chart(watch, ["recentAmount", "yoy", "sixMonthYoy"]))
    if not watch:
        watch_page += '<div class="notice">当前尚未记录观察类目的处理理由；完成商品层核验前不作入局判断。</div>'
    all_leaves = sorted((item for item in stats if item["level"] == 3),
                        key=lambda item: (item.get("growthImpactIndex") is None,
                                          -(item.get("growthImpactIndex") or 0), -item["recentAmount"], item["path"]))
    branch_options = ''.join(f'<option value="{esc(item["path"][-1])}">{esc(item["path"][-1])}</option>'
                             for item in stats if item["level"] == 2)
    all_head = ["二级类目", "三级类目", "前12个月", "最近12个月", "全年同比", "年度增量", "增长月份", "近6个月同比", "持续增量贡献", "状态",
                "容量路线", "容量百分位", "持续增长强度", "路线内排名", "半年增速差·百分点", "正增量前三月占比", "未入池原因"]
    all_rows = []
    for item in all_leaves:
        complete = item["complete"]
        raw_values = [item["path"][1], item["path"][-1], item["priorAmount"] if complete else None,
                      item["recentAmount"] if complete else None, item["yoy"], item["increment"],
                      item["positiveYoyMonths"], item["sixMonthYoy"], item.get("growthImpactIndex"),
                      "完整" if complete else "缺月"]
        shown = [item["path"][1], item["path"][-1], money(item["priorAmount"]) if complete else "—",
                 money(item["recentAmount"]) if complete else "—", percent(item["yoy"]),
                 money(item["increment"]), f'{item["positiveYoyMonths"]}/12' if complete else "—",
                 percent(item["sixMonthYoy"]), index_points(item.get("growthImpactIndex")), raw_values[-1]]
        status = raw_values[-1]
        raw_values += [item["capacityLabel"], item["capacityPercentile"], item["sustainedGrowthIndex"],
                       item["routeRank"], item["growthAccelerationPp"], item["top3PositiveIncrementShare"],
                       "；".join(item["routeExclusionReasons"]) or "已入池"]
        shown += [item["capacityLabel"], f'{item["capacityPercentile"]:.1f}' if item["capacityPercentile"] is not None else "—",
                  index_points(item["sustainedGrowthIndex"]), item["routeRank"] or "—",
                  f'{item["growthAccelerationPp"]:+.1f}' if item["growthAccelerationPp"] is not None else "—",
                  percent(item["top3PositiveIncrementShare"]), raw_values[-1]]
        cells = ''.join(f'<td data-sort="{esc(value if value is not None else "")}">{esc(label)}</td>'
                        for value, label in zip(raw_values, shown))
        all_rows.append(f'<tr data-second="{esc(item["path"][1])}" data-status="{esc(status)}" data-route="{esc(item["capacityBand"])}">{cells}</tr>')
    all_data = (f'<article class="card"><div class="card-head"><h2>全部三级类目</h2><span class="muted" id="allTableCount">{len(all_leaves)} 行</span></div>'
                f'<div class="table-controls"><label>二级类目 <select id="allSecondFilter"><option value="">全部</option>{branch_options}</select></label>'
                '<label>数据状态 <select id="allStatusFilter"><option value="">全部</option><option value="完整">完整</option><option value="缺月">缺月</option></select></label>'
                + '<label>容量路线 <select id="allRouteFilter"><option value="">全部</option>'
                + ''.join(f'<option value="{key}">{esc(label)}</option>' for key, label in LABELS.items()) + '</select></label>'
                +
                '<label>表内搜索 <input id="allTextFilter" type="search" placeholder="搜索三级类目"></label></div>'
                f'<div class="table-wrap"><table class="data-table sortable-table" id="allCategoryTable"{chart_attribute(stats_chart(all_leaves, ["priorAmount", "recentAmount", "yoy", "sixMonthYoy", "increment", "positiveYoyMonths", "growthImpactIndex", "sustainedGrowthIndex", "capacityPercentile", "routeRank", "growthAccelerationPp", "top3PositiveIncrementShare"]))}><thead><tr>'
                + ''.join(f'<th scope="col"><button type="button" data-sort-column="{index}" aria-label="按{esc(label)}排序">{esc(label)} <span aria-hidden="true">↕</span></button></th>'
                          for index, label in enumerate(all_head))
                + '</tr></thead><tbody>' + ''.join(all_rows) + '</tbody></table></div></article>')
    all_data += card("缺月范围", f'<p class="prose">{len(incomplete)} 个类目未覆盖全部24个月，完整年度同比不适用。逐项缺月见分析 JSON。</p>')
    method = card("统计口径", f'<ul class="list"><li>窗口：{esc(months[0])}—{esc(months[-1])}，前后各12个完整自然月。</li><li>同比＝近12个月成交额 ÷ 前12个月成交额 − 1；近6个月同比使用对应的上年6个月。</li><li>同月增长数＝后12个月中，成交额超过上年同月的月份数。</li><li>持续增量贡献（点）＝100 × max(年度成交增量, 0) ÷ 一级类目最近12个月成交额 × 同月增长数 ÷ 12；年度成交增量＝前12个月成交额 × 全年同比。缺月类目不计算。它仅用于本一级类目内部的深研排序，不能衡量白牌可进入性。</li></ul>')
    reconciliation_note = ("已通过可计算月份的层级金额核对。" if not issues
                           else "需复核类目映射或采集完整性。")
    method += route_method(route_model)
    method += card("下一步验证", '<ol class="list"><li>先轻量核查商品与店铺榜单的品牌集中度、非头部商家成交和价格带，注明样本覆盖；再决定深研顺序。</li><li>结合搜索词、评价和问答，识别细分场景与未满足需求。</li><li>核算采购、包装、物流、投放、退货和售后成本，再考虑小规模测试。</li></ol>')
    method += card("数据质量与证据边界", f'<ul class="list"><li>缺月类目：{len(incomplete)} 个；缺月类目不计算完整年度同比。</li><li>父子金额差异超过容差：{len(issues)} 处；{reconciliation_note}</li><li>24个月仅支持一个完整年度同比比较，不能证明多年持续增长。</li><li>尚缺品牌集中度、商品形态、消费者评价、采购成本、物流售后及投放费用。</li></ul>')
    source_items = ''.join(f'<li>{esc(item["file"])} · {esc(item["kind"])} · SHA-256 {esc(item["sha256"][:12])}…</li>' for item in provenance)
    method += card("数据来源", f'<ul class="list">{source_items}</ul>')
    pages = [
        ("overview", "一页看懂", view("overview", "决策摘要", f"{category}：哪些类目值得深研", "结论先行，后续页面逐层展示依据。", overview, show_heading=False)),
        ("structure", "市场结构", view("structure", "一级与二级", "市场增量来自哪里", "同口径比较二级类目规模、增量和占比。", structure)),
        ("candidates", "三条研究路线", view("candidates", "候选类目", "分路线选择自己的市场", "各路线独立排序；打开类目可查看完整逐月证据。", candidate)),
        ("watch", "观察与暂缓", view("watch", "反面证据", "体量、增速和可经营性不能互相替代", "观察不等于永久排除；等待新的商品和经营证据。", watch_page)),
        ("all", "全量类目", view("all", "全部数据", "三级类目可筛选、排序与比较", "统一展示所有三级类目，缺月项目单独标明。", all_data)),
        ("method", "口径与边界", view("method", "研究边界", "验证白牌差异化后再考虑试做", "市场规模和增长是研究信号。", method)),
    ]
    detail_pages = []
    detail_items = [(item, False) for item in candidates]
    if benchmark and benchmark not in candidates:
        detail_items.append((benchmark, True))
    for index, (item, is_benchmark) in enumerate(detail_items, start=1):
        detail_id = f"candidate-{index}"
        detail_pages.append((detail_id, item["path"][-1], view(
            detail_id, "规模基准" if is_benchmark else "首批深研", item["path"][-1],
            "基于类目月度数据的逐项解读；商品层假设另待验证。",
            candidate_detail(item, months, benchmark=is_benchmark))))
    nav_parts = []
    for name, label, _ in pages:
        if name != "candidates":
            nav_parts.append(f'<button data-view-target="{name}"><span class="nav-dot"></span>{esc(label)}</button>')
            continue
        nav_parts.append('<div class="nav-group" id="candidateNavGroup">'
                         '<div class="nav-group-header">'
                         f'<button data-view-target="{name}" data-nav-parent="candidates"><span class="nav-dot"></span>{esc(label)}</button>'
                         '<button type="button" id="candidateToggle" class="nav-disclosure" aria-controls="candidateSubmenu" aria-expanded="false" aria-label="展开研究路线目录"><span class="nav-chevron" aria-hidden="true">⌄</span></button></div>'
                         '<div class="nav-submenu" id="candidateSubmenu" hidden>')
        for detail_id, detail_label, _ in detail_pages:
            position = int(detail_id.split('-')[-1]) - 1
            item, is_benchmark = detail_items[position]
            prefix = "参照" if is_benchmark else item["capacityLabel"].replace("市场增量型", "").replace("容量持续增长型", "").replace("容量成长型", "")
            short_label = detail_label.split(" > ")[-1]
            nav_parts.append(f'<button class="nav-child" data-view-target="{detail_id}" title="{esc(detail_label)}">{esc(prefix)} · {esc(short_label)}</button>')
        nav_parts.append('</div></div>')
    nav = ''.join(nav_parts)
    replacements = {
        "__TITLE__": esc(f"{category}类目研究优先级报告"),
        "__NAV__": nav,
        "__SIDE_NOTE__": esc("24个月月度数据 · 一级至三级 · 仅供研究排序"),
        "__CATEGORY__": esc(category),
        "__PERIOD__": esc(f"{months[0]}—{months[-1]}"),
        "__CONTENT__": ''.join(page for _, _, page in pages + detail_pages),
        "__FOOTER__": esc("数据用于确定下一轮研究顺序，不构成备货或进入决定。"),
        "__CHART_SCRIPT__": (Path(__file__).resolve().parent.parent / "assets" / "chart-views.js").read_text(encoding="utf-8"),
    }
    output = template
    for key, value in replacements.items():
        if output.count(key) != 1:
            raise ValueError(f"报告模板缺少唯一占位符：{key}")
        output = output.replace(key, value)
    return output


def main():
    parser = argparse.ArgumentParser(description="Prepare an Excel cache or build a Yuce category opportunity HTML report")
    parser.add_argument("--category", required=True)
    parser.add_argument("--start-month", required=True)
    parser.add_argument("--end-month", required=True)
    parser.add_argument("--db-snapshot")
    parser.add_argument("--yc-json", action="append", default=[])
    parser.add_argument("--excel", action="append", default=[], help="Original Yuce category monthly .xlsx; repeat for all levels")
    parser.add_argument("--cache-xlsx", help="Previously prepared local Excel cache; analysis reads this, not the database")
    parser.add_argument("--cache-out", help="Create one new local data-only .xlsx cache from source inputs and exit")
    parser.add_argument("--decisions", help="Optional reviewed research decisions JSON")
    parser.add_argument("--strategy-config", help="Reviewed relative capacity quantiles and growth rules JSON")
    parser.add_argument("--analysis-only", action="store_true")
    parser.add_argument("--analysis-out")
    parser.add_argument("--html-out")
    args = parser.parse_args()
    months = months_between(args.start_month, args.end_month)
    if args.decisions:
        raise ValueError("v2 多路线请使用 --strategy-config 分位配置；旧 --decisions 绝对容量与单一名单须迁移")
    if args.analysis_only and args.html_out:
        raise ValueError("--analysis-only 不能同时提供 --html-out")
    if args.cache_out and (args.cache_xlsx or args.analysis_out or args.html_out or args.analysis_only or args.strategy_config):
        raise ValueError("生成缓存时只提供原始来源和 --cache-out")
    if not args.cache_out and not args.analysis_out:
        raise ValueError("生成分析需要 --analysis-out")
    if not args.cache_out and not args.analysis_only and not args.html_out:
        raise ValueError("生成报告需要 --html-out")
    analysis_out = Path(args.analysis_out) if args.analysis_out else None
    html_out = Path(args.html_out) if args.html_out else None
    cache_out = Path(args.cache_out) if args.cache_out else None
    for output in (analysis_out, html_out, cache_out):
        if output is None:
            continue
        if output.exists():
            raise ValueError(f"输出文件已存在，不能覆盖：{output}")
        if not output.parent.is_dir():
            raise ValueError(f"输出目录不存在：{output.parent}")
    if cache_out and cache_out.suffix.lower() != ".xlsx":
        raise ValueError("缓存输出必须是 .xlsx")
    merged, provenance = merge_rows(args.category, months, args.db_snapshot, args.yc_json,
                                    args.excel, args.cache_xlsx)
    stats = compute_stats(args.category, months, merged)
    issues = reconcile(months, merged)
    if cache_out:
        write_cache(cache_out, merged)
        print(json.dumps({"ok": True, "cache": str(cache_out), "rowCount": len(merged),
                          "categoryCount": len(stats), "sourceCount": len(provenance),
                          "reconciliationIssueCount": len(issues)}, ensure_ascii=False))
        return
    supplied_policy = read_json(args.strategy_config)[0] if args.strategy_config else None
    route_model = build_routes(stats, supplied_policy)
    index = {tuple(x["path"]): x for x in stats if x["level"] == 3}
    pool = [index[tuple(path)] for route in route_model["routes"] for path in route["poolPaths"]]
    candidates = [index[tuple(path)] for route in route_model["routes"] for path in route["candidatePaths"]]
    benchmark = max((x for x in index.values() if x["complete"]), key=lambda x: x["recentAmount"], default=None)
    screen = {"minRecentAmount": route_model["boundaryAmounts"][0] or 0, **route_model["policy"]}
    watch = quantitative_watch(stats, candidates, benchmark, screen)
    decision_mode = "reviewed-relative-routes" if args.strategy_config else "provisional-relative-routes"
    analysis = {"contract": CONTRACT, "category": args.category, "months": months,
                "sources": provenance, "stats": stats, "reconciliationIssues": issues,
                "marketRoutes": route_model,
                "screeningPoolPaths": [item["path"] for item in pool],
                "candidatePaths": [item["path"] for item in candidates],
                "benchmarkPath": benchmark["path"] if benchmark else None,
                "watchPaths": [item["path"] for item in watch],
                "decisionMode": decision_mode}
    report = None
    if not args.analysis_only:
        template = (Path(__file__).resolve().parent.parent / "assets" / "report-template.html").read_text(encoding="utf-8")
        report = render_html(args.category, months, stats, pool, candidates, benchmark, watch, screen,
                             issues, provenance, template, decision_mode, route_model)
    # Exclusive creation keeps every prior research run recoverable.
    with analysis_out.open("x", encoding="utf-8") as target:
        json.dump(analysis, target, ensure_ascii=False, indent=2)
        target.write("\n")
    if html_out:
        with html_out.open("x", encoding="utf-8") as target:
            target.write(report)
    print(json.dumps({"ok": True, "contract": CONTRACT, "analysis": str(analysis_out),
                      "html": str(html_out) if html_out else None, "categoryCount": len(stats),
                      "candidateCount": len(candidates), "reconciliationIssueCount": len(issues)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
