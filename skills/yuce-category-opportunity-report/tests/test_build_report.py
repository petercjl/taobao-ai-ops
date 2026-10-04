import json
import re
from html import unescape
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook

SKILL = Path(__file__).resolve().parents[1]
SCRIPT = SKILL / "scripts" / "build_report.py"


def months():
    values = []
    year, month = 2024, 9
    for _ in range(24):
        values.append(f"{year:04d}-{month:02d}")
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    return values


class BuildReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.period = months()

    def write(self, name, content):
        target = self.root / name
        target.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        return target

    def sources(self):
        first = self.write("first.json", {"kind": "first-category-monthly", "category": "测试类目", "months": self.period,
            "rows": [{"month": month, "firstCategory": "测试类目", "transactionAmount": 1000 if index < 12 else 1300}
                     for index, month in enumerate(self.period)]})
        second = self.write("second.json", {"kind": "second-category-monthly", "category": "测试类目", "months": self.period,
            "rows": [{"month": month, "firstCategory": "测试类目", "secondCategory": "二级A", "transactionAmount": 1000 if index < 12 else 1300,
                      "hasThirdCategory": "是"} for index, month in enumerate(self.period)]})
        third = self.write("third.json", {"kind": "third-category-monthly", "category": "测试类目", "secondCategory": "二级A", "months": self.period,
            "rows": [{"month": month, "firstCategory": "测试类目", "secondCategory": "二级A", "thirdCategory": "三级A",
                      "transactionAmount": 1000 if index < 12 else 1300} for index, month in enumerate(self.period)]})
        return first, second, third

    def excel_source(self):
        target = self.root / "monthly.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "一级类目月数据"
        first.append(["月份", "一级类目", "成交额（原始值）", "成交量（原始值）"])
        second = workbook.create_sheet("二级类目月数据")
        second.append(["月份", "一级类目", "二级类目", "总成交额（原始值）", "是否有三级类目"])
        third = workbook.create_sheet("三级类目月数据")
        third.append(["月份", "一级类目", "二级类目", "三级类目", "总成交额（原始值）"])
        for index, month in enumerate(self.period):
            amount = 1000 if index < 12 else 1300
            first.append([month, "测试类目", amount, 100])
            second.append([month, "测试类目", "二级A", amount, "是"])
            third.append([month, "测试类目", "二级A", "三级A", amount])
        workbook.save(target)
        workbook.close()
        return target

    def run_report(self, sources):
        command = [sys.executable, str(SCRIPT), "--category", "测试类目", "--start-month", self.period[0],
                   "--end-month", self.period[-1], "--analysis-out", str(self.root / "analysis.json"),
                   "--html-out", str(self.root / "report.html")]
        for source in sources:
            command.extend(["--yc-json", str(source)])
        return subprocess.run(command, capture_output=True, text=True)

    def test_builds_complete_report_and_refuses_overwrite(self):
        result = self.run_report(self.sources())
        self.assertEqual(result.returncode, 0, result.stderr)
        analysis = json.loads((self.root / "analysis.json").read_text(encoding="utf-8"))
        report = (self.root / "report.html").read_text(encoding="utf-8")
        tables = re.findall(r'<table\b([^>]*)>(.*?)</table>', report, re.S)
        self.assertTrue(tables)
        for attrs, contents in tables:
            encoded = re.search(r'data-chart-spec="([^"]*)"', attrs)
            self.assertIsNotNone(encoded, "Every table must have a chart spec")
            spec = json.loads(unescape(encoded.group(1)))
            tbody = re.search(r'<tbody>(.*?)</tbody>', contents, re.S).group(1)
            self.assertEqual(len(spec["labels"]), len(re.findall(r'<tr\b', tbody)))
            for group in spec["groups"]:
                for series in group["series"]:
                    self.assertEqual(len(series["values"]), len(spec["labels"]))
        self.assertNotIn("__CHART_SCRIPT__", report)
        self.assertEqual(analysis["candidatePaths"], [["测试类目", "二级A", "三级A"]])
        self.assertEqual(analysis["reconciliationIssues"], [])
        self.assertIn("一页看懂", report)
        self.assertIn("这份数据最重要的三句话", report)
        self.assertIn("先研究哪里", report)
        self.assertIn("三条研究路线", report)
        self.assertIn('id="allRouteFilter"', report)
        self.assertEqual(report.count('class="view" data-view='), 7)
        self.assertIn('id="view-watch"', report)
        self.assertIn('class="nav-child" data-view-target="candidate-1"', report)
        self.assertIn('id="allCategoryTable"', report)
        self.assertIn('id="allSecondFilter"', report)
        self.assertIn('data-sort-column="8"', report)
        self.assertIn('一级及二级类目逐月成交额', report)
        self.assertIn('12组同月成交额对比', report)
        self.assertEqual(len(analysis["stats"][0]["monthlyAmounts"]), 24)
        self.assertEqual(len(analysis["screeningPoolPaths"]), 1)
        leaf = next(item for item in analysis["stats"] if item["level"] == 3)
        self.assertAlmostEqual(leaf["growthImpactIndex"], 100 * 3600 / 15600, places=3)
        self.assertEqual(leaf["routeRank"], 1)
        self.assertEqual(self.run_report(self.sources()).returncode, 1)

    def test_conflicting_sources_stop_before_output(self):
        first, second, third = self.sources()
        changed = json.loads(third.read_text(encoding="utf-8"))
        changed["rows"][0]["transactionAmount"] = 999
        conflict = self.write("conflict.json", changed)
        result = self.run_report([first, second, third, conflict])
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("来源数值冲突", result.stderr)
        self.assertFalse((self.root / "report.html").exists())

    def test_zero_market_renders_three_empty_routes(self):
        sources = self.sources()
        for source in sources:
            data = json.loads(source.read_text())
            for row in data["rows"]:
                row["transactionAmount"] = 0
            source.write_text(json.dumps(data))
        result = self.run_report(sources)
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads((self.root / "analysis.json").read_text())
        self.assertEqual([r["poolCount"] for r in data["marketRoutes"]["routes"]], [0, 0, 0])
        self.assertEqual(data["candidatePaths"], [])
        report = (self.root / "report.html").read_text()
        for label in ("大市场增量型", "中等容量持续增长型", "中小容量成长型"):
            self.assertIn(label, report)

    def test_no_third_level_still_renders_empty_routes(self):
        first, second, _ = self.sources()
        result = self.run_report([first, second])
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads((self.root / "analysis.json").read_text())
        self.assertEqual(data["candidatePaths"], [])
        self.assertEqual(len(data["marketRoutes"]["routes"]), 3)

    def test_absolute_decisions_require_migration(self):
        result = subprocess.run([sys.executable, str(SCRIPT), "--category", "测试类目",
            "--start-month", self.period[0], "--end-month", self.period[-1],
            "--decisions", "legacy.json", "--analysis-only", "--analysis-out", str(self.root / "no.json")],
            capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("须迁移", result.stderr)
        self.assertFalse((self.root / "no.json").exists())

    def test_reviewed_relative_policy_is_saved(self):
        source = self.excel_source()
        config = self.write("policy.json", {"capacityQuantiles": [0.3, 0.65, 0.85],
            "why": "根据当前测试总体的分布与经营目标调整相对分位位置。"})
        result = subprocess.run([sys.executable, str(SCRIPT), "--category", "测试类目",
            "--start-month", self.period[0], "--end-month", self.period[-1],
            "--excel", str(source), "--strategy-config", str(config),
            "--analysis-out", str(self.root / "analysis.json"),
            "--html-out", str(self.root / "report.html")], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads((self.root / "analysis.json").read_text())
        self.assertEqual(data["marketRoutes"]["policy"]["capacityQuantiles"], [0.3, 0.65, 0.85])
        self.assertEqual(len(data["marketRoutes"]["routes"]), 3)
        self.assertEqual(data["decisionMode"], "reviewed-relative-routes")

    def test_excel_input_creates_data_only_cache_and_report(self):
        source = self.excel_source()
        cache = self.root / "cache.xlsx"
        base = [sys.executable, str(SCRIPT), "--category", "测试类目",
                "--start-month", self.period[0], "--end-month", self.period[-1]]
        prepared = subprocess.run(base + ["--excel", str(source), "--cache-out", str(cache)],
                                  capture_output=True, text=True)
        self.assertEqual(prepared.returncode, 0, prepared.stderr)
        workbook = load_workbook(cache, read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["类目月度缓存"])
        rows = list(workbook.active.values)
        self.assertEqual(len(rows) - 1, 72)
        self.assertEqual({row[1] for row in rows[1:]}, {1, 2, 3})
        self.assertEqual(len({row[0] for row in rows[1:]}), 24)
        workbook.close()
        self.assertNotEqual(subprocess.run(base + ["--excel", str(source), "--cache-out", str(cache)],
                                          capture_output=True, text=True).returncode, 0)
        analysis = self.root / "cache-analysis.json"
        report = self.root / "cache-report.html"
        rendered = subprocess.run(base + ["--cache-xlsx", str(cache), "--analysis-out", str(analysis),
                                          "--html-out", str(report)], capture_output=True, text=True)
        self.assertEqual(rendered.returncode, 0, rendered.stderr)
        content = json.loads(analysis.read_text(encoding="utf-8"))
        self.assertEqual(content["candidatePaths"], [["测试类目", "二级A", "三级A"]])
        self.assertTrue(any(item["file"] == "monthly.xlsx" for item in content["sources"]))

    def test_excel_amount_conflict_stops_before_cache(self):
        source = self.excel_source()
        conflict = self.root / "changed.xlsx"
        workbook = load_workbook(source)
        workbook["三级类目月数据"]["E2"] = 999
        workbook.save(conflict)
        workbook.close()
        cache = self.root / "cache.xlsx"
        result = subprocess.run([sys.executable, str(SCRIPT), "--category", "测试类目",
                                 "--start-month", self.period[0], "--end-month", self.period[-1],
                                 "--excel", str(source), "--excel", str(conflict), "--cache-out", str(cache)],
                                capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("来源数值冲突", result.stderr)
        self.assertFalse(cache.exists())


if __name__ == "__main__":
    unittest.main()
